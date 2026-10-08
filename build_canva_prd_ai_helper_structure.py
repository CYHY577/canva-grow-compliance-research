from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\pc\Desktop\canva测试")
ASSETS = ROOT / "prd-assets"
OUTDIR = ROOT / "output"
OUTDIR.mkdir(parents=True, exist_ok=True)
OUT = OUTDIR / "Canva_Grow_营销内容合规层_PRD_AI工具助手结构版.docx"
FLOW = ASSETS / "product_flow_swimlane.png"

PURPLE = "7D2AE8"
PURPLE_DARK = "5B20B5"
INK = "222127"
MUTED = "68636D"
LIGHT = "F7F5F9"
BORDER = "DED9E4"
GREEN = "167A53"
GREEN_BG = "EAF7F1"
AMBER = "946500"
AMBER_BG = "FFF4D6"
RED = "B73B2D"
RED_BG = "FFF0ED"
BLUE = "2563C7"
BLUE_BG = "EDF4FF"
GRAY_BG = "F0EFF2"
FONT = "Microsoft YaHei"


def make_flowchart():
    w, h = 2400, 1760
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    font_path = r"C:\Windows\Fonts\msyh.ttc"
    bold_path = r"C:\Windows\Fonts\msyhbd.ttc"
    f_title = ImageFont.truetype(bold_path, 52)
    f_lane = ImageFont.truetype(bold_path, 34)
    f_box = ImageFont.truetype(font_path, 28)
    f_small = ImageFont.truetype(font_path, 24)
    f_label = ImageFont.truetype(bold_path, 24)
    d.text((80, 45), "Canva AI 海报文字合规审核：产品上线后泳道流程", fill="#222127", font=f_title)
    d.text((80, 112), "审核对象是最终海报中的可见文字；Prompt 仅用于生成，不进入合规判断。", fill="#68636D", font=f_small)

    x0, y0 = 70, 180
    lane_label_w = 250
    lane_h = 350
    lane_names = ["用户", "Canva AI", "Compliance\nLayer", "知识与\n规则服务"]
    lane_fills = ["#F9F7FB", "#F4EEFC", "#EEF5FF", "#F2F8F5"]
    for i, name in enumerate(lane_names):
        y = y0 + i * lane_h
        d.rounded_rectangle((x0, y, w-70, y+lane_h-10), radius=22, fill=lane_fills[i], outline="#D8D3DE", width=3)
        d.rounded_rectangle((x0, y, x0+lane_label_w, y+lane_h-10), radius=22, fill="#FFFFFF", outline="#D8D3DE", width=3)
        lines=name.split("\n")
        for j,line in enumerate(lines):
            lf=f_lane if len(line)<=8 else f_box
            bbox = d.textbbox((0,0), line, font=lf)
            d.text((x0+(lane_label_w-(bbox[2]-bbox[0]))/2, y+130+j*44), line, fill="#5B20B5", font=lf)

    cols = [390, 760, 1130, 1500, 1870]
    box_w, box_h = 280, 118

    def box(lane, col, text, fill="#FFFFFF", stroke="#9D91A8", bold=False, hh=None):
        yy = y0 + lane*lane_h + 105
        xx = cols[col]
        bh = hh or box_h
        d.rounded_rectangle((xx, yy, xx+box_w, yy+bh), radius=20, fill=fill, outline=stroke, width=4)
        ff = f_label if bold else f_box
        lines = text.split("\n")
        total = len(lines)*38
        for j, line in enumerate(lines):
            bb = d.textbbox((0,0), line, font=ff)
            d.text((xx+(box_w-(bb[2]-bb[0]))/2, yy+(bh-total)/2+j*38), line, fill="#222127", font=ff)
        return (xx, yy, xx+box_w, yy+bh)

    def arrow(a, b, color="#7D2AE8", label=None, dotted=False):
        ax, ay = (a[2], (a[1]+a[3])//2)
        bx, by = (b[0], (b[1]+b[3])//2)
        if abs(by-ay) < 10:
            pts = [(ax+8, ay), (bx-12, by)]
        else:
            mx = (ax+bx)//2
            pts = [(ax+8, ay), (mx, ay), (mx, by), (bx-12, by)]
        if dotted:
            for k in range(len(pts)-1):
                p1,p2=pts[k],pts[k+1]
                steps=max(1,int(((p2[0]-p1[0])**2+(p2[1]-p1[1])**2)**0.5//22))
                for s in range(0,steps,2):
                    t1=s/steps; t2=min((s+1)/steps,1)
                    d.line((p1[0]+(p2[0]-p1[0])*t1,p1[1]+(p2[1]-p1[1])*t1,p1[0]+(p2[0]-p1[0])*t2,p1[1]+(p2[1]-p1[1])*t2),fill=color,width=5)
        else:
            d.line(pts, fill=color, width=6, joint="curve")
        d.polygon([(bx-12,by),(bx-30,by-11),(bx-30,by+11)],fill=color)
        if label:
            d.text(((ax+bx)//2-55, min(ay,by)-33), label, fill=color, font=f_small)

    def down(a, b, color="#7D2AE8", label=None):
        ax=(a[0]+a[2])//2; ay=a[3]+8; bx=(b[0]+b[2])//2; by=b[1]-12
        mx=(ax+bx)//2
        d.line([(ax,ay),(ax,(ay+by)//2),(bx,(ay+by)//2),(bx,by)],fill=color,width=6,joint="curve")
        d.polygon([(bx,by),(bx-11,by-18),(bx+11,by-18)],fill=color)
        if label: d.text((mx+10,(ay+by)//2-34),label,fill=color,font=f_small)

    def route_up_right(a, b, color="#7D2AE8", label=None):
        ax=a[2]+8; ay=(a[1]+a[3])//2; bx=b[2]+8; by=(b[1]+b[3])//2
        edge=w-115
        d.line([(ax,ay),(edge,ay),(edge,by),(bx,by)],fill=color,width=6,joint="curve")
        d.polygon([(bx,by),(bx+18,by-11),(bx+18,by+11)],fill=color)
        if label: d.text((edge-10, (ay+by)//2),label,fill=color,font=f_small)

    u1 = box(0,0,"输入海报\n生成需求",fill="#FFFFFF",stroke="#7D2AE8",bold=True)
    ai1 = box(1,0,"生成背景图\n与文字图层",fill="#F7F0FF",stroke="#7D2AE8")
    ai2 = box(1,1,"组合为可用\n营销海报",fill="#F7F0FF",stroke="#7D2AE8")
    c1 = box(2,1,"渲染完成\n自动触发",fill="#EDF4FF",stroke="#2B6DE5",bold=True)
    c2 = box(2,2,"读取文字图层\nOCR 兜底",fill="#EDF4FF",stroke="#2B6DE5")
    k1 = box(3,2,"校验 Context\n路由知识包",fill="#EDF8F2",stroke="#17865A")
    c3 = box(2,3,"规则与模型\n协同判断",fill="#EDF4FF",stroke="#2B6DE5")
    c4 = box(2,4,"返回四态\n结果与下一步",fill="#F7F0FF",stroke="#7D2AE8",bold=True)
    u2 = box(0,4,"查看结果\n执行下一步",fill="#FFFFFF",stroke="#7D2AE8",bold=True)

    down(u1, ai1, label="生成")
    arrow(ai1, ai2)
    down(ai2, c1, label="完成事件")
    arrow(c1, c2)
    down(c2, k1, label="Context")
    arrow(k1, c3, color="#17865A", label="已发布规则")
    arrow(c3, c4)
    route_up_right(c4, u2, label="呈现")

    # action band
    band_y = 1560
    d.rounded_rectangle((300, band_y, w-80, h-55), radius=24, fill="#FAF8FC", outline="#D8D3DE", width=3)
    states = [
        ("可以发布", "继续/下载/发布", "#EAF7F1", "#167A53"),
        ("发布前需修改", "修改后重检", "#FFF0ED", "#B73B2D"),
        ("需要补充信息", "补 Context 后重检", "#FFF4D6", "#946500"),
        ("不在覆盖范围", "查看覆盖边界", "#F0EFF2", "#68636D"),
    ]
    for i,(name,act,fill,stroke) in enumerate(states):
        xx=335+i*500
        d.rounded_rectangle((xx,band_y+28,xx+450,h-65),radius=18,fill=fill,outline=stroke,width=4)
        d.text((xx+24,band_y+48),name,fill=stroke,font=f_label)
        d.text((xx+24,band_y+86),act,fill="#333038",font=f_small)
    d.text((80, 1490), "结果变化规则：文字、Context 或规则版本任一变化 → 旧结果失效 → 自动重检；应用改写本身不等于通过。", fill="#5B20B5", font=f_label)
    im.save(FLOW, quality=95)


def shade(cell, fill):
    pr=cell._tc.get_or_add_tcPr(); sh=pr.find(qn("w:shd"))
    if sh is None: sh=OxmlElement("w:shd"); pr.append(sh)
    sh.set(qn("w:fill"),fill)


def margins(cell, top=110, start=130, bottom=110, end=130):
    pr=cell._tc.get_or_add_tcPr(); mar=pr.first_child_found_in("w:tcMar")
    if mar is None: mar=OxmlElement("w:tcMar"); pr.append(mar)
    for name,val in (("top",top),("start",start),("bottom",bottom),("end",end)):
        e=mar.find(qn(f"w:{name}"))
        if e is None: e=OxmlElement(f"w:{name}"); mar.append(e)
        e.set(qn("w:w"),str(val)); e.set(qn("w:type"),"dxa")


def geometry(table,widths,indent=0):
    pr=table._tbl.tblPr
    tw=pr.find(qn("w:tblW"))
    if tw is None: tw=OxmlElement("w:tblW"); pr.append(tw)
    tw.set(qn("w:w"),str(sum(widths))); tw.set(qn("w:type"),"dxa")
    ti=pr.find(qn("w:tblInd"))
    if ti is None: ti=OxmlElement("w:tblInd"); pr.append(ti)
    ti.set(qn("w:w"),str(indent)); ti.set(qn("w:type"),"dxa")
    lay=pr.find(qn("w:tblLayout"))
    if lay is None: lay=OxmlElement("w:tblLayout"); pr.append(lay)
    lay.set(qn("w:type"),"fixed")
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col=OxmlElement("w:gridCol"); col.set(qn("w:w"),str(width)); grid.append(col)
    for row in table.rows:
        for i,cell in enumerate(row.cells):
            prc=cell._tc.get_or_add_tcPr(); tcw=prc.find(qn("w:tcW"))
            if tcw is None: tcw=OxmlElement("w:tcW"); prc.append(tcw)
            tcw.set(qn("w:w"),str(widths[i])); tcw.set(qn("w:type"),"dxa")
            margins(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    pr=row._tr.get_or_add_trPr(); e=OxmlElement("w:tblHeader"); e.set(qn("w:val"),"true"); pr.append(e)


def runfmt(run,size=10.5,bold=False,color=INK,italic=False):
    run.font.name=FONT
    rpr=run._element.get_or_add_rPr(); rpr.rFonts.set(qn("w:ascii"),FONT); rpr.rFonts.set(qn("w:hAnsi"),FONT); rpr.rFonts.set(qn("w:eastAsia"),FONT)
    run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=RGBColor.from_string(color)
    return run


def para(p,before=0,after=5,line=1.15,keep=False):
    f=p.paragraph_format; f.space_before=Pt(before); f.space_after=Pt(after); f.line_spacing=line; f.keep_with_next=keep
    return p


def body(doc,text,after=5,color=INK,bold_prefix=None):
    p=doc.add_paragraph(); para(p,after=after)
    if bold_prefix and text.startswith(bold_prefix):
        runfmt(p.add_run(bold_prefix),bold=True,color=color); runfmt(p.add_run(text[len(bold_prefix):]),color=color)
    else: runfmt(p.add_run(text),color=color)
    return p


def bullet(doc,text,level=0):
    p=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2")
    p.paragraph_format.left_indent=Inches(0.48 if level==0 else 0.72); p.paragraph_format.first_line_indent=Inches(-0.24); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.12
    runfmt(p.add_run(text),size=10.2); return p


def number(doc,text):
    p=doc.add_paragraph(style="List Number"); p.paragraph_format.left_indent=Inches(.48); p.paragraph_format.first_line_indent=Inches(-.24); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.12
    runfmt(p.add_run(text),size=10.2); return p


def heading(doc,text,level=1):
    p=doc.add_paragraph(text,style=f"Heading {level}"); p.paragraph_format.keep_with_next=True; return p


def kicker(doc,text):
    p=doc.add_paragraph(); para(p,before=12,after=3,keep=True); runfmt(p.add_run(text.upper()),size=8.8,bold=True,color=PURPLE); return p


def callout(doc,label,text,fill="F5F0FC",color=PURPLE):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.LEFT; geometry(t,[9410],0)
    c=t.cell(0,0); shade(c,fill); p=c.paragraphs[0]; para(p,after=0,line=1.18)
    runfmt(p.add_run(label+"  "),size=10.2,bold=True,color=color); runfmt(p.add_run(text),size=10.2)
    sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(0); return t


def table(doc,headers,rows,widths,font_size=8.8,header_fill="F0EDF4"):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.LEFT
    repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        shade(t.rows[0].cells[i],header_fill); p=t.rows[0].cells[i].paragraphs[0]; para(p,after=0,line=1.05); runfmt(p.add_run(h),size=font_size,bold=True)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            if ri%2: shade(cells[i],"FBFAFC")
            p=cells[i].paragraphs[0]; para(p,after=0,line=1.12); runfmt(p.add_run(str(val)),size=font_size)
    geometry(t,widths,0); sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(0); return t


def image(doc,filename,caption,width=6.45):
    pth=ASSETS/filename
    if not pth.exists(): body(doc,f"[截图缺失：{filename}]",color=RED); return
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(3)
    p.add_run().add_picture(str(pth),width=Inches(width))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(8); runfmt(c.add_run(caption),size=8.2,color=MUTED,italic=True)


def page_number(p):
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    runfmt(p.add_run("第 "),size=8,color=MUTED)
    r=p.add_run()._r; b=OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"),"begin"); ins=OxmlElement("w:instrText"); ins.set(qn("xml:space"),"preserve"); ins.text=" PAGE "; e=OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"),"end"); r.extend([b,ins,e])
    runfmt(p.add_run(" 页"),size=8,color=MUTED)


make_flowchart()
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(1.75); sec.bottom_margin=Cm(1.75); sec.left_margin=Cm(2.15); sec.right_margin=Cm(2.15); sec.header_distance=Cm(0.85); sec.footer_distance=Cm(0.85)

normal=doc.styles["Normal"]; normal.font.name=FONT; normal._element.rPr.rFonts.set(qn("w:eastAsia"),FONT); normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(INK); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.15
for level,size,before,after,color in [(1,16,14,7,INK),(2,13,10,5,PURPLE_DARK),(3,11.3,7,3,INK)]:
    st=doc.styles[f"Heading {level}"]; st.font.name=FONT; st._element.rPr.rFonts.set(qn("w:eastAsia"),FONT); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True

header=sec.header; ht=header.add_table(rows=1,cols=2,width=Inches(6.54)); geometry(ht,[5600,3810],0)
runfmt(ht.cell(0,0).paragraphs[0].add_run("CANVA GROW · PRODUCT REQUIREMENTS DOCUMENT"),size=7.8,bold=True,color=MUTED)
ht.cell(0,1).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT; runfmt(ht.cell(0,1).paragraphs[0].add_run("AU · BEAUTY · POSTER TEXT"),size=7.8,bold=True,color=PURPLE)
page_number(sec.footer.paragraphs[0])

# Cover
p=doc.add_paragraph(); para(p,before=18,after=4); runfmt(p.add_run("PRODUCT REQUIREMENTS DOCUMENT"),size=9.2,bold=True,color=PURPLE)
p=doc.add_paragraph(); para(p,after=6); runfmt(p.add_run("Canva Grow 营销内容合规层 PRD"),size=24,bold=True)
p=doc.add_paragraph(); para(p,after=18); runfmt(p.add_run("基于《AI工具助手PRD》结构重构 · 图片/海报生成后的可见文字审核"),size=12.5,color=MUTED)
callout(doc,"一句话定位","用户在 Canva AI 生成可使用的广告图片或海报后，系统自动读取最终画面中的可见文字，并结合国家、行业、产品监管身份与受控知识包，返回明确状态和下一步操作。")
table(doc,["字段","内容"],[
    ["产品名称","Canva Grow Compliance Layer"],
    ["本期范围","Australia · Beauty · Image / Poster · Visible text"],
    ["主要用户","营销创作者、品牌人员、中小企业主；运营侧为法规与知识运营人员"],
    ["版本日期","V2.0 · 2026-09-03"],
    ["Demo 基线","Canva AI · Australia Compliance Layer 最新网页"],
],[1800,7610],font_size=9.2)
kicker(doc,"KEY PRODUCT PRINCIPLE")
heading(doc,"系统审查的是成品海报，不是用户输入的 Prompt",1)
body(doc,"Prompt 仅用于图片与文案生成。合规判断必须以最终渲染后的海报可见文字为准：优先读取 Canva 原生文字图层；对于上传或扁平化图片，使用 OCR 兜底。")
callout(doc,"价值边界","系统协助用户完成判断，不替用户做法律裁决；“可以发布”仅代表在已识别、已覆盖、已检查的范围内未发现明确问题。",fill="EDF4FF",color=BLUE)

doc.add_page_break()
kicker(doc,"00 · DOCUMENT")
heading(doc,"0. 文档信息",1)
table(doc,["项目","说明"],[
    ["产品定位","嵌入 Canva AI 图片生成与使用链路的营销内容合规决策层。"],
    ["核心任务","在用户真正准备使用营销海报时，自动扫描画面文字，判断信息是否充分、可能适用何种规则、具体哪里需处理以及下一步怎么做。"],
    ["本期能力","图片生成、可见文字提取、Context 采集、知识包路由、四态判断、最小改写、自动重检、发布前 Gate。"],
    ["明确不做","Prompt 合规审查、视觉主体判断、商标判断、平台政策、落地页、视频和音频。"],
],[1900,7510],font_size=9.2)

kicker(doc,"01 · BACKGROUND")
heading(doc,"1. 项目背景",1)
heading(doc,"1.1 业务背景",2)
body(doc,"Canva AI 已能帮助用户快速生成广告图片和海报，但“生成完成”不等于“可以直接使用”。在美妆等受监管行业中，最终画面里的功效声称、强制声明、证据依赖和产品监管身份会共同影响使用决策。")
bullet(doc,"营销用户通常不熟悉法域、产品监管分类和强制披露要求。")
bullet(doc,"问题可能是“少说了什么”，而不是出现明显违禁词，例如 AUST L 产品缺少必要声明。")
bullet(doc,"合规介入过早会打断创意探索，介入过晚则造成返工和发布风险。")
bullet(doc,"因此，产品需要把检查放在“海报已经形成、即将被编辑或使用”的节点。")
heading(doc,"1.2 用户痛点",2)
table(doc,["痛点","现状","需要的产品能力"],[
    ["判断信息不足","用户不知道还缺产品监管身份或证据","系统先判断信息是否充分，缺失时进入“需要补充信息”"],
    ["规则难以理解","用户无法从法规中快速定位与当前文案相关的义务","展示可能适用的规则、知识包版本和覆盖边界"],
    ["问题难以处理","只给风险提示，用户仍不知道改哪里","定位海报文字区域并给最小改写建议"],
    ["容易被误拦截","主观偏好或成分事实也可能被泛化判高风险","提供 Checked and fine，证明系统看过但不拦"],
],[1800,3200,4410],font_size=8.9)

kicker(doc,"02 · WHY")
heading(doc,"2. 为什么做",1)
heading(doc,"2.1 产品机会",2)
body(doc,"与独立合规工具相比，Canva 的优势是拥有生成、文字图层、编辑和导出上下文，可以把判断直接嵌入创作链路。用户无需复制文案到外部工具，也无需猜测该在什么时候运行检查。")
callout(doc,"核心链路","自然语言需求 → 生成海报 → 渲染完成 → 提取最终可见文字 → 路由知识包 → 返回四态结果 → 修改/补充/继续使用。")
heading(doc,"2.2 用户需求优先级",2)
table(doc,["维度","判断方式","本项目结论"],[
    ["用户痛点强度","是否高频影响发布效率","Beauty 广告中的功效声称、强制披露和证据问题影响明显"],
    ["业务价值","是否影响生成到使用的转化","减少末端返工，提高有效发布完成率"],
    ["AI 适配度","是否需要语义理解和区域定位","适合由 OCR/文字图层、规则和语言模型协同完成"],
    ["实施及风险成本","能否建立可控范围","P0 限定 AU · Beauty · Poster text，并使用版本化知识包"],
],[1700,3200,4510],font_size=8.9)
heading(doc,"2.3 AI 可行性验证",2)
bullet(doc,"图像生成有效性：用户输入后必须生成新图片，不得静默沿用示例素材。")
bullet(doc,"文字提取准确性：原生文字图层优先；OCR 需返回区域坐标与置信度。")
bullet(doc,"规则路由准确性：Country、Industry、Product Regulation、Format 能命中正确知识包。")
bullet(doc,"决策一致性：相同文字、Context 与规则版本应返回可复现结果。")
bullet(doc,"异常兜底：知识覆盖不足、OCR 失败或模型超时时不得输出绿色结论。")

kicker(doc,"03 · GOALS")
heading(doc,"3. 产品目标",1)
body(doc,"帮助用户在不离开 Canva AI 创作链路的情况下，理解海报文字在当前覆盖范围内能否继续使用，以及下一步应修改、补充信息还是继续发布。")
number(doc,"在可用营销海报形成后自动介入，不打断概念探索。")
number(doc,"正确识别最终海报中的文字，并将问题定位到具体区域。")
number(doc,"将法规知识、产品监管身份、广告表达与 supporting evidence 组合成可解释判断。")
number(doc,"用四种 Action Status 替代笼统风险分级。")
number(doc,"内容、Context 或规则版本变化后自动失效并重检。")
number(doc,"每次结果明确展示已检查与未检查范围。")
heading(doc,"3.1 成功指标",2)
table(doc,["指标","定义","建议试点目标"],[
    ["有效发布完成率","进入检查后，以有效 Ready 状态完成下载/发布/继续的会话占比","建立基线后提升"],
    ["关键遗漏召回率","强制声明或明确禁止项被识别的比例","≥95%"],
    ["误拦截率","无需修改的海报文字被判 Fix 的比例","≤8%"],
    ["可见文字提取成功率","文字层或 OCR 形成可审核文本的比例","≥98%"],
    ["覆盖披露率","每次结果展示已检查/未检查的比例","100%"],
    ["P95 决策耗时","文字稳定到结果可用，不含图片生成时间","≤4 秒"],
],[2200,4900,2310],font_size=8.8)
body(doc,"注：目标值为试点门槛建议，需要通过影子模式、专家标注和真实流量校准。",color=MUTED)

doc.add_page_break()
kicker(doc,"04 · CORE FEATURES")
heading(doc,"4. 核心功能设计",1)
heading(doc,"4.1 海报生成与审核对象",2)
body(doc,"用户在 Canva AI 输入广告需求后，系统生成背景图和可编辑文字层，并组合为可使用海报。Prompt 只进入生成链路，不作为合规判断依据。")
table(doc,["对象","用途","是否参与审核"],[
    ["用户 Prompt","指导图片与文案生成","否"],
    ["背景图","海报视觉主体","P0 不判断视觉合规"],
    ["Canva 文字图层","品牌名、标题、正文、声明和脚注","是，优先读取"],
    ["扁平化图片文字","上传或烘焙在图片中的文字","是，OCR 兜底"],
    ["Context","国家、行业、产品、监管身份、渠道、格式","是，用于知识路由"],
],[1900,3500,4010],font_size=9.0)
image(doc,"00-generation-to-check.png","图 1  最新网页：用户输入需求后生成新的广告海报，系统随后扫描最终画面文字。")

heading(doc,"4.2 Canva 内 AI 文案审核的触发方式与入口（新增）",2)
body(doc,"审核入口与 Canva AI 生成结果绑定，不新建独立法务页面。海报渲染完成后，状态条紧贴该条生成结果的下沿；在侧边形态中，同一份结果显示在右侧抽屉，用户可边看边修改。")
table(doc,["触发事件","触发条件","系统动作","防止的问题"],[
    ["首次生成","图片/海报渲染完成，存在文字层或 OCR 结果","自动提取文字、确认 Context 并检查","不在 Prompt 输入阶段提前打断"],
    ["文字实质修改","用户停止编辑或确认编辑","旧结果失效，debounce 后重检","避免逐字调用"],
    ["Context 变化","国家、行业、产品、监管身份、渠道或格式变化","重新路由知识包并重检","避免沿用旧规则"],
    ["应用建议改写","建议写回文字层","重新渲染并执行完整检查","改写不等于通过"],
    ["准备继续使用","下载/发布/继续时结果不存在、过期或失败","执行 Pre-use Compliance Gate","避免异常路径绕过"],
],[1400,2650,2900,2460],font_size=8.3)
bullet(doc,"概念探索阶段不运行完整检查；只有可实际使用的营销内容形成后才介入。")
bullet(doc,"Content Hash + Context + Rule Version 完全一致时复用最近一次有效结果。")
bullet(doc,"无对应知识包时返回“不在覆盖范围”，不得让大模型自行补规则。")
image(doc,"05-inline-entry.png","图 2  页面内入口：窄状态条紧贴生成结果下沿，固定展示状态、主操作与覆盖边界。")

heading(doc,"4.3 Action Status 与推进权限",2)
table(doc,["状态","系统判断","用户下一步","下载/发布/继续"],[
    ["可以发布","在已识别、已覆盖的文字范围内未发现明确问题","查看详情或继续","允许"],
    ["发布前需修改","存在明确问题或缺少强制声明","修改或应用建议后重检","阻止"],
    ["需要补充信息","判断依赖缺失的 Context 或证据","补充最少必要信息后重检","阻止"],
    ["不在覆盖范围","当前国家/行业/产品/格式没有已发布知识包","查看覆盖范围或转人工","阻止"],
    ["暂时无法检查","OCR、规则或模型服务失败","重试或转人工","阻止"],
],[2000,3100,2850,1460],font_size=8.6)
callout(doc,"措辞要求","“可以发布”必须带“在已检查范围内”；禁止出现“完全合规”“已通过”“安全”或“保证”等绝对表达。",fill=GREEN_BG,color=GREEN)

heading(doc,"4.4 结果详情与处理动作",2)
bullet(doc,"被标记的海报文字区域与原文片段。")
bullet(doc,"问题类型与一句话原因。")
bullet(doc,"可能适用的规则，不写死未经确认的条款号。")
bullet(doc,"最小改写建议，并明确改动内容。")
bullet(doc,"已检查/未检查两列、知识包版本与 trace_id。")
bullet(doc,"用户修改或应用建议后必须自动重检；修复动作本身不算通过。")

heading(doc,"4.5 Context 采集与使用引导",2)
table(doc,["字段","P0 示例","缺失处理"],[
    ["Country","Australia","必填"],
    ["Industry","Beauty","必填"],
    ["Product Type","Skincare","无法判断时追问"],
    ["Product Regulation","Cosmetic / AUST L / Not sure","Not sure 且含功效声称时进入需要补充信息"],
    ["Channel","Instagram","P0 仅用于 Context；平台政策未检查"],
    ["Content Format","Image / Poster","非支持格式返回不在覆盖范围"],
],[2100,3200,4010],font_size=8.8)

heading(doc,"4.6 异常及兜底",2)
table(doc,["场景","处理原则","用户可见结果"],[
    ["OCR 低置信或无文字","不把不确定文字当作完整输入","要求确认文字或显示暂时无法检查"],
    ["知识包缺失/过期","禁止模型生成新规则","不在覆盖范围"],
    ["模型/规则服务超时","旧结果立即失效","暂时无法检查，可重试"],
    ["图片生成失败","不得保留示例图冒充结果","明确生成失败，可重新生成"],
    ["内容未变化","复用最近一次有效结果","不重复运行"],
],[2100,4300,3010],font_size=8.8)

doc.add_page_break()
kicker(doc,"05 · PRODUCT FLOW")
heading(doc,"5. 产品整体流程",1)
body(doc,"以下流程只描述产品发布后的真实用户路径，不包含评审用展示形式切换。四条泳道分别是用户、Canva AI、Compliance Layer 和知识与规则服务。")
image(doc,"product_flow_swimlane.png","图 3  产品上线后泳道流程：从海报生成到自动审核、规则判断和用户下一步。",width=6.55)
heading(doc,"5.1 主流程说明",2)
number(doc,"用户输入海报需求；Prompt 仅用于生成。")
number(doc,"Canva AI 生成背景图与可编辑文字层，组合为可用海报。")
number(doc,"渲染完成事件自动触发审核，提取文字层；扁平图使用 OCR。")
number(doc,"系统校验 Context 并路由已发布知识包。")
number(doc,"规则层处理强制义务和覆盖，模型负责语义理解、区域定位和受控改写。")
number(doc,"前端返回四态结果；非有效 Ready 阻止下载、发布和继续，但不阻断编辑与对话。")
number(doc,"文字、Context 或规则版本变化时，旧结果失效并重检。")

kicker(doc,"06 · CASES")
heading(doc,"6. 典型案例与页面对应",1)
heading(doc,"6.1 Case A：AUST L 缺少强制声明",2)
body(doc,"场景：海报没有明显违禁词，但 Product Regulation 为 AUST L Listed Medicine，最终画面缺少 “ALWAYS READ THE LABEL AND FOLLOW THE DIRECTIONS FOR USE”。")
body(doc,"结果：发布前需修改。系统定位必要声明区域，阻止下载/发布/继续；补充声明后重新渲染并完整重检。")
image(doc,"01-trigger-fix-side-wide.png","图 4  发布前需修改：指出缺失声明、可能规则、最小改写和覆盖边界。")

heading(doc,"6.2 Case B：产品监管身份不确定",2)
body(doc,"场景：海报含量化功效声称，但 Product Regulation 为 Not sure。仅靠文字无法确定应按 Cosmetic 还是 AUST L 规则判断。")
body(doc,"结果：需要补充信息。系统要求用户确认监管身份；如保留量化或临床声称，还需提供与当前产品和范围匹配的 supporting evidence。")
image(doc,"02-needs-input-side.png","图 5  需要补充信息：缺失必要 Context 时不猜测、不输出绿色结论。")

heading(doc,"6.3 Case C：系统看过但无需拦截",2)
body(doc,"场景：海报包含“我们团队最爱的一款配方”、成分事实和非量化肤感描述。")
body(doc,"结果：可以发布。详情中展示 Checked and fine，并说明“看过这句，判定无需修改”，证明系统既能发现问题，也能避免误拦截。")
image(doc,"03-ready-checked-fine.png","图 6  可以发布：结论限定在已识别并检查的海报文字范围内。")

kicker(doc,"07 · SCOPE")
heading(doc,"7. 功能范围",1)
heading(doc,"7.1 P0 — 本期",2)
bullet(doc,"Canva AI 广告图片/海报生成；背景图与可编辑文字层组合。")
bullet(doc,"渲染完成、文字修改、Context 变化、应用改写和继续使用前自动触发。")
bullet(doc,"原生文字图层读取与 OCR 兜底。")
bullet(doc,"Australia · Beauty · Image/Poster · Visible text 知识包。")
bullet(doc,"Cosmetic、AUST L Listed Medicine 与 Not sure 路径。")
bullet(doc,"四态结果、问题定位、最小改写、覆盖披露、版本与 trace_id。")
bullet(doc,"自动失效、自动重检、结果缓存与 Pre-use Compliance Gate。")
heading(doc,"7.2 P1",2)
bullet(doc,"上传/扁平化图片的增强 OCR；更多 AU Beauty 产品类型。")
bullet(doc,"平台政策知识包与 Landing Page 文本联动。")
bullet(doc,"支持更多渠道的差异化规则。")
heading(doc,"7.3 P2",2)
bullet(doc,"视频关键帧与字幕、音频转写、多页面设计。")
bullet(doc,"更多国家、行业和语言的独立知识包。")

kicker(doc,"08 · FUTURE")
heading(doc,"8. 后续能力",1)
table(doc,["方向","能力","前置条件"],[
    ["多模态扩展","视觉主体、人物、包装和品牌资产检查","建立专用视觉规则和标注集"],
    ["跨内容联动","海报、落地页和平台政策联合判断","统一内容 ID 与多包路由"],
    ["证据工作台","将功效声称关联到企业 supporting evidence","权限、保留期和证据版本治理"],
    ["知识运营自动化","来源更新提醒、影响分析和候选规则生成","专家审核后才允许发布"],
],[1800,4100,3510],font_size=8.9)

kicker(doc,"09 · EVALUATION")
heading(doc,"9. 评测",1)
body(doc,"评测重点不是模型是否能泛泛谈法规，而是面对真实海报与真实 Context 时，能否稳定提取文字、命中正确知识包、做出正确 Action Status，并返回可执行下一步。")
heading(doc,"9.1 评测集构成",2)
table(doc,["样本类型","示例","关键判断"],[
    ["强制披露缺失","AUST L 海报缺必要声明","零违禁词也必须判 Fix"],
    ["监管边界","Cosmetic 文案出现 repair damaged skin","识别可能跨越化妆品边界"],
    ["证据依赖","clinically proven / 30% in 14 days","证据缺失或不匹配时 Need input/Fix"],
    ["Hard negative","团队最爱、成分事实、非量化肤感","应 Ready，不误拦"],
    ["OCR 压力","小字、低对比、复杂背景、弯曲文字","低置信不得假装完整"],
    ["范围外","非 AU、非 Beauty 或非海报格式","Outside scope"],
],[1800,4200,3410],font_size=8.8)
heading(doc,"9.2 核心指标",2)
bullet(doc,"文字提取字符准确率与区域召回率。")
bullet(doc,"知识包路由准确率与规则 ID 正确率。")
bullet(doc,"四态准确率、关键遗漏召回率与严重漏放率。")
bullet(doc,"误拦截率、修复后完成率和用户理解度。")
bullet(doc,"P95 延迟、超时率、版本一致性与可回滚性。")
callout(doc,"发布门禁","严重漏放必须为 0；边界样本不得输出 Ready；证据与规则映射需 100% 可追溯。",fill=RED_BG,color=RED)

kicker(doc,"10 · SYSTEM")
heading(doc,"10. 系统与数据要求",1)
table(doc,["组件","职责","禁止事项"],[
    ["CogView 图像生成","根据 Prompt 生成广告背景图","不输出合规结论"],
    ["GLM 文案生成","生成品牌名、标题和可编辑文字层","不决定适用法规"],
    ["文字提取层","读取 Canva 文字图层，OCR 识别扁平图","低置信文字不得当作确定输入"],
    ["知识路由/规则层","确定覆盖、强制义务、版本和 Gate","知识包缺失时不得让模型补规则"],
    ["GLM 合规推理","语义理解、片段定位、解释和最小改写","不得发明来源、条款或规则 ID"],
],[2000,4100,3310],font_size=8.7)
heading(doc,"10.1 决策输出契约",2)
table(doc,["字段","要求"],[
    ["status","READY / FIX / NEED_INPUT / OUTSIDE_SCOPE / ERROR"],
    ["extraction","source、regions、confidence、text_hash"],
    ["issues","region_id、span、issue_type、reason、suggestion"],
    ["coverage","checked、not_checked、jurisdiction、format"],
    ["knowledge_version","如 AU-BEA-POSTER-TEXT v1.4"],
    ["models","copy_model、image_model、decision_model"],
    ["trace_id / latency","全链路唯一 ID 与分阶段耗时"],
],[2400,7010],font_size=9.0)
heading(doc,"10.2 安全与隐私",2)
bullet(doc,"API Key 仅保存在服务端 Secret，前端和日志不得出现明文。")
bullet(doc,"默认不使用客户 Prompt、海报或原文训练模型；进入评测集需授权并去标识。")
bullet(doc,"生产日志优先记录内容哈希、区域坐标、状态、规则 ID 和版本。")
bullet(doc,"临时图片 URL 需有有效期，生产环境转存到受控存储。")

kicker(doc,"11 · ACCEPTANCE")
heading(doc,"11. P0 验收标准",1)
table(doc,["场景","Given / When","Then"],[
    ["首次生成","用户输入新海报需求","生成新图片与文字层；渲染完成后自动审核，不沿用示例内容"],
    ["缺强制声明","AUST L 海报缺 mandatory statement","返回 Fix，定位问题；补充后重检才允许继续"],
    ["监管身份未知","Not sure 且海报含功效声称","返回 Need input；不得根据 Prompt 猜测分类"],
    ["检查无问题","AU Cosmetic 海报文字无明确问题","返回 Ready，并展示 Checked and fine、覆盖和版本"],
    ["文字修改","用户停止编辑海报文字","旧结果立即失效，debounce 后自动重检"],
    ["应用改写","建议写回文字层","重新渲染并完整重检；改写动作本身不算通过"],
    ["范围不支持","非 AU Beauty Poster 组合","返回 Outside scope，不显示绿色结论"],
    ["服务异常","OCR、模型或知识包不可用","返回 Error，阻止关键推进并提供重试"],
],[1800,3400,4210],font_size=8.6)
heading(doc,"11.1 上线阶段",2)
table(doc,["阶段","范围","退出条件"],[
    ["Shadow","后台运行，不阻止用户","完成基线、遗漏和误拦截校准"],
    ["Assisted","展示四态；仅高置信强制义务硬阻断","专家复核稳定，用户理解覆盖边界"],
    ["Controlled GA","白名单 AU Beauty 用户与知识包版本","指标达标、回滚演练完成"],
    ["Expansion","扩展更多产品、格式和规则包","每个包独立通过评测门禁"],
],[1800,4000,3610],font_size=8.8)

kicker(doc,"APPENDIX")
heading(doc,"附录：P0 功能清单",1)
table(doc,["ID","功能","验收摘要"],[
    ["F01","海报生成","输入后生成新背景图和文字层"],
    ["F02","自动触发","渲染完成、修改、Context 变化、应用改写、继续前"],
    ["F03","文字提取","文字图层优先，OCR 兜底并返回区域/置信度"],
    ["F04","Context","六字段可追溯，缺失时追问"],
    ["F05","知识路由","仅命中已发布且有效的 AU 知识包"],
    ["F06","四态决策","Ready / Fix / Need input / Outside scope；异常为 Error"],
    ["F07","问题定位","区域、原文、原因、可能规则和下一步"],
    ["F08","最小改写","写回后自动重检"],
    ["F09","覆盖披露","每次展示 checked/not checked、版本和 trace_id"],
    ["F10","发布 Gate","非有效 Ready 阻止下载、发布和继续"],
    ["F11","缓存失效","文字/Context/版本任一变化即失效"],
    ["F12","安全密钥","模型 Key 仅存服务端 Secret"],
],[700,2300,6410],font_size=8.7)

doc.core_properties.title="Canva Grow 营销内容合规层 PRD（AI工具助手结构版）"
doc.core_properties.subject="Canva AI poster visible-text compliance review"
doc.core_properties.author="AI Product · Risk & Compliance"
doc.core_properties.keywords="Canva, AI, Compliance Layer, Poster, OCR, Australia, Beauty"
doc.save(OUT)
print(OUT)
