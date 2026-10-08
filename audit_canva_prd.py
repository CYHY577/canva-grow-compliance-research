from pathlib import Path
from zipfile import ZipFile
from docx import Document
from lxml import etree

p = Path(r"C:\Users\pc\Desktop\canva测试\output\Canva_Grow_营销内容合规层_PRD_飞书导入版.docx")
with ZipFile(p) as z:
    print("zip_test", z.testzip())
    xml = z.read("word/document.xml")
    print("manual_page_breaks", xml.count(b'w:type="page"'))
    print("tables", xml.count(b"<w:tbl>"))
    print("images", xml.count(b"<a:blip"))

d = Document(p)
text = " ".join(v.text for v in d.paragraphs)
text += " " + " ".join(c.text for t in d.tables for row in t.rows for c in row.cells)
print("paragraphs", len(d.paragraphs))
print("sections", len(d.sections))
print("inline_shapes", len(d.inline_shapes))
print("bad_tokens", [s for s in ["TODO", "TBD", "turn0", "codex-file-citation", "PLACEHOLDER"] if s in text])
print("size_bytes", p.stat().st_size)

with ZipFile(p) as z:
    root = etree.fromstring(z.read("word/document.xml"))
ns = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
page = 1
stats = {page: {"paras": 0, "tables": 0, "rows": 0, "image_height_in": 0.0}}
for child in root.find("w:body", ns):
    if child.tag.endswith("}p"):
        stats[page]["paras"] += 1
        for ext in child.xpath(".//a:ext", namespaces=ns):
            stats[page]["image_height_in"] += int(ext.get("cy", 0)) / 914400
        if child.xpath('.//w:br[@w:type="page"]', namespaces=ns):
            page += 1
            stats[page] = {"paras": 0, "tables": 0, "rows": 0, "image_height_in": 0.0}
    elif child.tag.endswith("}tbl"):
        stats[page]["tables"] += 1
        stats[page]["rows"] += len(child.xpath("./w:tr", namespaces=ns))
        for ext in child.xpath(".//a:ext", namespaces=ns):
            stats[page]["image_height_in"] += int(ext.get("cy", 0)) / 914400
print("page_stats")
for n, s in stats.items():
    print(n, s)
