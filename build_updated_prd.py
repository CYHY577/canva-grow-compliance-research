from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\pc\Desktop\canva测试")
ASSETS = ROOT / "prd-assets"
OUTDIR = ROOT / "output"
OUTDIR.mkdir(parents=True, exist_ok=True)
OUT = OUTDIR / "Canva_Grow_营销内容合规层_PRD_更新版.docx"

PURPLE = "7D2AE8"
PURPLE_DARK = "5B20B5"
BLUE = "2B6DE5"
CYAN = "00AEB8"
GREEN = "17865A"
GREEN_BG = "EAF8F2"
AMBER = "AD7600"
AMBER_BG = "FFF5D9"
RED = "C43B27"
RED_BG = "FFF0EC"
INK = "25212B"
MUTED = "6F6875"
LIGHT = "F6F4F8"
BORDER = "DDD8E2"
WHITE = "FFFFFF"
FONT = "Microsoft YaHei"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=11, bold=False, color=INK, italic=False, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_para(p, before=0, after=6, line=1.1, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep
    return p


def add_body(doc, text, bold_prefix=None, after=6, color=INK):
    p = doc.add_paragraph()
    set_para(p, after=after, line=1.15)
    if bold_prefix and text.startswith(bold_prefix):
        set_run(p.add_run(bold_prefix), bold=True, color=color)
        set_run(p.add_run(text[len(bold_prefix):]), color=color)
    else:
        set_run(p.add_run(text), color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    pf = p.paragraph_format
    pf.left_indent = Inches(0.5 if level == 0 else 0.75)
    pf.first_line_indent = Inches(-0.25)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.15
    set_run(p.add_run(text), size=10.7)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    pf = p.paragraph_format
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.25)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.15
    set_run(p.add_run(text), size=10.7)
    return p


def add_kicker(doc, text):
    p = doc.add_paragraph()
    set_para(p, before=16, after=5, keep=True)
    set_run(p.add_run(text.upper()), size=9.3, bold=True, color=PURPLE)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def callout(doc, label, text, color=PURPLE, fill="F5F0FC"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360], indent=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_para(p, after=0, line=1.2)
    set_run(p.add_run(label + "  "), size=10.5, bold=True, color=color)
    set_run(p.add_run(text), size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, widths, header_fill=LIGHT, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        set_cell_shading(hdr.cells[i], header_fill)
        p = hdr.cells[i].paragraphs[0]
        set_para(p, after=0, line=1.05)
        set_run(p.add_run(value), size=font_size, bold=True, color=INK)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "FBFAFC")
            p = cells[i].paragraphs[0]
            set_para(p, after=0, line=1.15)
            set_run(p.add_run(str(value)), size=font_size, color=INK)
    set_table_geometry(table, widths, indent=120)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    return table


def add_image(doc, filename, caption, width=6.35):
    path = ASSETS / filename
    if not path.exists():
        add_body(doc, f"[截图缺失：{filename}]", color=RED)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(9)
    set_run(c.add_run(caption), size=8.5, color=MUTED, italic=True)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r = paragraph.add_run()._r
    r.append(fld_char1)
    r.append(instr)
    r.append(fld_char2)
    set_run(paragraph.add_run(" 页"), size=8.5, color=MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.42)
section.footer_distance = Inches(0.42)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for level, size, before, after, color in [
    (1, 16, 16, 8, PURPLE_DARK),
    (2, 13, 12, 6, PURPLE),
    (3, 11.5, 8, 4, "3C3542"),
]:
    st = styles[f"Heading {level}"]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

header = section.header
ht = header.add_table(rows=1, cols=2, width=Inches(6.5))
set_table_geometry(ht, [5200, 4160], indent=0)
ht.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
ht.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run(ht.cell(0, 0).paragraphs[0].add_run("CANVA GROW · COMPLIANCE LAYER PRD"), size=8, bold=True, color=MUTED)
set_run(ht.cell(0, 1).paragraphs[0].add_run("AU · BEAUTY · POSTER TEXT"), size=8, bold=True, color=PURPLE)
footer = section.footer
fp = footer.paragraphs[0]
add_page_number(fp)

# Cover / executive summary
p = doc.add_paragraph()
set_para(p, before=10, after=4)
set_run(p.add_run("PRODUCT REQUIREMENTS DOCUMENT"), size=9.5, bold=True, color=PURPLE)
p = doc.add_paragraph()
set_para(p, after=5)
set_run(p.add_run("Canva Grow 营销内容合规层 PRD"), size=25, bold=True, color=INK)
p = doc.add_paragraph()
set_para(p, after=18)
set_run(p.add_run("图片/海报生成后的可见文字审核 · Australia Beauty P0"), size=13, color=MUTED)

add_kicker(doc, "PROJECT VALUE")
heading(doc, "项目对 Canva 的价值", 1)
value_table = doc.add_table(rows=2, cols=4)
value_table.style = "Table Grid"
value_rows = [
    [("1", "更快发布", "创作后即时检查\n减少末端返工", "E8FBFC", CYAN),
     ("2", "行业增长", "从 Beauty 起步\n扩展高监管场景", "F5EEFF", PURPLE),
     ("3", "可信 AI", "可解释、可追溯\n明确覆盖边界", "EAF8F2", GREEN),
     ("4", "知识复利", "评测驱动迭代\n形成规则资产", "FFF6DD", AMBER)],
]
for i, item in enumerate(value_rows[0]):
    num, title, desc, fill, color = item
    c = value_table.cell(0, i)
    set_cell_shading(c, fill)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, after=2)
    set_run(p.add_run(num + "  " + title), size=11, bold=True, color=color)
    d = c.add_paragraph()
    d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(d, after=0, line=1.2)
    set_run(d.add_run(desc), size=9, color=MUTED)
for i, (label, value) in enumerate([
    ("发布转化", "减少生成后返工"),
    ("合规成本", "把检查前移"),
    ("覆盖能力", "图片文字可判断"),
    ("知识治理", "版本化与可回滚"),
]):
    c = value_table.cell(1, i)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, after=1)
    set_run(p.add_run(label), size=9.5, bold=True, color=INK)
    d = c.add_paragraph()
    d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(d, after=0)
    set_run(d.add_run(value), size=8.7, color=MUTED)
set_table_geometry(value_table, [2340, 2340, 2340, 2340], indent=120)

callout(doc, "核心判断", "这不是检测用户 Prompt 的工具。系统只在广告图片或海报完成生成后，读取最终画面中的可见文字，再结合国家、行业、产品监管身份和知识包协助用户完成合规判断。", PURPLE, "F4EEFC")
add_table(doc, ["字段", "内容"], [
    ["Owner", "AI Product · Risk & Compliance"],
    ["状态", "Updated for product alignment · 2026-09-02"],
    ["P0 范围", "Australia · Beauty · Image / Poster · Visible text"],
    ["产品定位", "嵌入 Canva AI 图片生成到发布链路的合规决策层"],
], [1800, 7560], header_fill="F1EEF5", font_size=9.5)

add_kicker(doc, "01 · CONTEXT")
heading(doc, "1. 文档信息与项目背景", 1)
add_table(doc, ["字段", "内容"], [
    ["产品名称", "Canva Grow Compliance Layer（营销内容合规层）"],
    ["产品定位", "广告图片/海报生成完成后，自动扫描最终可见文字，给出可执行的判断、修复建议或信息追问。"],
    ["目标用户", "品牌营销人员、中小企业主、代理商创作者；运营端为法规专家、内容政策与风险团队。"],
    ["本期范围", "Australia · Beauty · Image / Poster；仅审核可见文字；支持 Cosmetic 与 AUST L Listed Medicine。"],
    ["明确边界", "不审核用户 Prompt；不判断人物、产品视觉、商标、平台政策或落地页；不替代法律意见。"],
], [1800, 7560], font_size=9.3)

heading(doc, "1.1 为什么现在做", 2)
add_table(doc, ["用户旅程", "当前断点", "产品机会"], [
    ["生成海报", "模型能生成图片，但用户不知道画面上的文案是否满足市场与品类义务。", "图片渲染完成后立即扫描文字并给出状态。"],
    ["编辑海报", "文字修改后原判断可能失效。", "监听文字图层与 OCR 结果变化，停止编辑后自动重检。"],
    ["准备发布", "合规常在末端介入，导致返工或等待人工。", "对非 Ready 状态禁用下载、发布与继续。"],
    ["跨市场复用", "用户易把局部规则理解为全市场覆盖。", "固定披露已检查/未检查和知识包版本。"],
], [1600, 3600, 4160], font_size=9.1)

heading(doc, "1.2 用户与 Jobs-to-be-done", 2)
add_table(doc, ["角色", "核心任务", "成功体验"], [
    ["创作者", "快速生成并使用广告图片/海报", "输入需求后获得新图片与文字；问题可直接处理，不离开当前对话。"],
    ["品牌/法务", "控制强制披露、功效声称和证据风险", "可追溯到文字区域、规则包、来源和 trace_id。"],
    ["知识运营", "维护市场与产品规则", "知识包有版本、有效期、评测和回滚，不依赖通用模型记忆。"],
], [1500, 3600, 4260], font_size=9.2)

add_kicker(doc, "02 · GOALS & SCOPE")
heading(doc, "2. 产品目标、成功指标与范围", 1)
heading(doc, "2.1 北极星与护栏指标", 2)
add_table(doc, ["指标", "定义", "试点建议目标", "用途"], [
    ["有效发布完成率", "进入检查后，以有效 Ready 状态完成发布/下载的会话占比", "建立基线后提升", "北极星"],
    ["关键遗漏召回率", "强制声明或明确禁用项被识别的比例", "≥95%", "安全护栏"],
    ["误拦截率", "真实可用海报文字被判 Fix 的比例", "≤8%", "体验护栏"],
    ["可见文字提取成功率", "文字图层或 OCR 可形成可审查文本的海报占比", "≥98%", "输入质量"],
    ["覆盖披露率", "每次结果展示 checked/not checked 的比例", "100%", "透明度"],
    ["P95 决策耗时", "文字稳定到结果可用；不含图片生成时间", "≤4s", "性能"],
    ["修复后完成率", "应用建议后通过重检并继续的比例", "≥60%", "价值验证"],
], [1900, 3760, 1900, 1800], font_size=8.8)
add_body(doc, "注：以上为 PRD 阶段建议阈值，需通过影子模式、专家标注与真实流量校准，不得作为现状数据对外传播。", color=MUTED)

heading(doc, "2.2 范围优先级", 2)
add_table(doc, ["阶段", "包含", "明确不包含"], [
    ["P0", "图片/海报生成；可编辑文字层；OCR 兜底；AU Beauty 文字审核；四态结果；修复、重检和发布 Gate。", "视觉主体、人物形象、商标、Meta 平台政策、落地页。"],
    ["P1", "上传/扁平化图片 OCR；更多 AU 品类；平台政策包；Landing Page 文本。", "自动扩展为所有市场。"],
    ["P2", "视频关键帧/字幕、音频转写、多市场自助接入。", "无专家审核的自动法规发布。"],
], [1000, 4860, 3500], font_size=9.1)
callout(doc, "范围变化", "原 PRD 的 Australia · Beauty · Text 已更新为 Australia · Beauty · Image / Poster · Visible text。图片是内容载体，但 P0 判断对象仍是最终画面中的文字。", BLUE, "EEF4FF")

add_kicker(doc, "03 · CORE EXPERIENCE")
heading(doc, "3. 核心体验：生成后扫描，而不是检查 Prompt", 1)
heading(doc, "3.1 输入、生成对象与审核对象", 2)
add_table(doc, ["对象", "系统用途", "是否进入合规判断"], [
    ["用户 Prompt", "指导广告图片和可编辑文字层生成", "否。不得根据 Prompt 推断产品分类或规则。"],
    ["生成背景图", "形成海报视觉主体", "P0 不判断视觉合规；要求尽量不烘焙文字。"],
    ["Canva 文字图层", "承载品牌名、标题、正文、必要声明", "是，优先读取。"],
    ["扁平化图片文字", "上传或栅格化后的画面文字", "是，使用 OCR 兜底并保留区域坐标与置信度。"],
    ["Context", "国家、行业、产品、监管身份、渠道、格式", "是，用于知识包路由和边界判断。"],
], [1800, 3700, 3860], font_size=9.1)
callout(doc, "产品原则", "Prompt 用于“生成”，海报可见文字用于“审核”。两条链路在产品和日志中必须显式分离。", PURPLE, "F4EEFC")

heading(doc, "3.2 Canva 内触发方式与入口（新增）", 2)
add_body(doc, "Compliance Check 的入口紧贴 Canva AI 生成结果，不建立独立法务页面。图片完成渲染后，状态条随生成结果一起出现在对话流中；在侧边形态下，同一份结果显示于右侧抽屉。")
add_table(doc, ["触发事件", "触发条件", "系统行为", "不应发生"], [
    ["首次生成", "广告图片/海报已完成渲染，存在可用文字层或 OCR 结果", "自动提取文字，确认 Context，运行检查", "输入 Prompt 时提前审查"],
    ["实质性编辑", "用户修改文字层并停止编辑/确认编辑", "旧结果标记失效，debounce 后重检", "逐字触发请求"],
    ["Context 变化", "国家、行业、产品、监管身份、渠道或格式变化", "重新路由知识包并重检", "复用旧规则包"],
    ["应用建议改写", "建议被写回文字层", "重新渲染并执行完整检查", "把“已改写”当作通过"],
    ["准备继续使用", "下载、发布、继续时结果不存在、过期、失败或规则版本失效", "触发 Pre-use Compliance Gate", "绕过检查"],
], [1400, 2750, 3050, 2160], font_size=8.5)
add_bullet(doc, "不在概念探索阶段运行完整检查；只有可实际使用的营销海报形成后才介入。")
add_bullet(doc, "内容哈希 + Context + Rule Version 与最近一次有效结果一致时，直接复用，不重复调用模型。")
add_bullet(doc, "不支持的组合返回 Outside scope，不得生成绿色结论。")
add_image(doc, "05-inline-entry.png", "图 1  页面内入口：合规状态条紧贴生成结果下沿，展示状态、覆盖披露与详情入口。")

heading(doc, "3.3 结果状态与权限", 2)
add_table(doc, ["状态", "触发条件", "用户动作", "推进权限"], [
    ["Ready to run / 可以发布", "Context 完整、范围支持、当前可见文字未发现明确问题", "查看详情后继续", "允许；文案必须带“在已检查范围内”"],
    ["Fix before running / 发布前需修改", "发现强制披露缺失或明确问题", "应用建议或手动编辑，再重检", "阻止发布/下载/继续"],
    ["Need your input / 需要补充信息", "产品监管身份、证据或必要 Context 不足", "补充最少必要信息", "阻止发布/下载/继续"],
    ["Outside scope / 不在覆盖范围", "国家、行业、产品或格式无已发布知识包", "调整范围或转人工", "不得出现绿色结论"],
    ["Error / 暂时无法检查", "提取、模型或规则服务失败", "重试或转人工", "旧结果不可继续使用"],
], [1900, 3300, 2700, 1460], font_size=8.8)

heading(doc, "3.4 信息架构与展示方式", 2)
add_bullet(doc, "收起态最多包含：状态徽章、一句话、主操作，以及固定的已检查/未检查说明。")
add_bullet(doc, "展开后显示：被标记文字区域、问题原因、可能适用的规则、最小改写、覆盖范围、知识包版本和 trace_id。")
add_bullet(doc, "侧边抽屉不接管画布，用户可边看提示边编辑；抽屉可折叠为 48px 状态窄条。")
add_bullet(doc, "后三种不可继续状态只阻止下载/发布/继续，不阻止用户继续对话和编辑。")

add_kicker(doc, "04 · END-TO-END FLOW")
heading(doc, "4. 端到端产品流程", 1)
for step in [
    "用户在 Canva AI 输入广告需求。Prompt 仅进入生成链路。",
    "服务端调用图像生成模型生成无文字或少文字的背景图，并由语言模型生成可编辑文字层。",
    "Canva 将背景图和文字图层组合为可实际使用的海报。",
    "渲染完成事件触发文字提取：原生文字图层优先，OCR 作为扁平化图片兜底。",
    "系统校验 Country、Industry、Product、Product Regulation、Channel、Content Format，并路由已发布知识包。",
    "规则层处理强制声明、覆盖范围和发布 Gate；模型负责语义理解、片段定位与解释改写。",
    "UI 返回四态结果；用户修改、补充信息或继续使用；任何变化均使旧结果失效并重检。",
]:
    add_number(doc, step)
add_image(doc, "00-generation-to-check.png", "图 2  真实生成链路：输入需求后生成新广告背景图与文字层，随后自动扫描最终海报文字并给出结果。")

heading(doc, "4.1 数据流与隔离要求", 2)
add_table(doc, ["阶段", "输入", "输出", "关键控制"], [
    ["Poster Generation", "用户 Prompt", "背景图 URL + 可编辑文字层", "Key 仅保存在服务端 Secret；Prompt 不传给合规模型。"],
    ["Text Extraction", "文字图层或图片", "文本区域、坐标、来源、OCR 置信度", "优先使用结构化文字图层；记录提取版本。"],
    ["Knowledge Routing", "Context + format", "适用知识包或 Outside scope", "不存在知识包时禁止模型补规则。"],
    ["Compliance Decision", "提取文字 + Context + 受控规则", "status、issues、recommendation、coverage", "允许 ID 白名单；强规则优先；输出可审计。"],
    ["Pre-use Gate", "当前结果状态与版本", "允许或阻止继续", "非有效 Ready 一律阻止关键推进。"],
], [1700, 2250, 2600, 2810], font_size=8.7)

add_kicker(doc, "05 · CASES")
heading(doc, "5. 核心案例与页面对应", 1)
heading(doc, "5.1 Case A：缺少强制声明 → Fix before running", 2)
add_body(doc, "Given：AUST L Listed Medicine 的海报已生成，但最终可见文字中不存在 “ALWAYS READ THE LABEL AND FOLLOW THE DIRECTIONS FOR USE”。")
add_body(doc, "Then：定位到底部必要声明区域，阻止下载、发布和继续；用户补充声明后必须重新渲染并完整重检。")
add_image(doc, "01-trigger-fix-side-wide.png", "图 3  发布前需修改：系统扫描最终海报，指出缺失内容、可能规则、最小改写与覆盖边界。")

heading(doc, "5.2 Case B：产品监管身份不确定 → Need your input", 2)
add_body(doc, "Given：海报含 “Clinically proven…30% in 14 days” 的量化功效声称，但 Product Regulation 为 Not sure。")
add_body(doc, "Then：系统不根据 Prompt 猜测分类，要求用户选择 Cosmetic 或 AUST L Listed Medicine；如保留量化声称，还需提供匹配的 supporting evidence。")
add_image(doc, "02-needs-input-side.png", "图 4  需要补充信息：缺失 Context 时不做法律猜测，也不输出绿色结论。")

heading(doc, "5.3 Case C：检查过但无需拦截 → Ready to run", 2)
add_body(doc, "Given：海报包含“我们团队最爱的一款配方”、成分事实和非量化肤感描述。")
add_body(doc, "Then：系统明确说明“看过这句，判定无需修改”，展示 Checked and fine，证明产品不仅能发现问题，也能避免误拦截。")
add_image(doc, "03-ready-checked-fine.png", "图 5  可以发布：状态限定为已识别并检查的海报文字范围，视觉元素、平台政策和落地页仍未检查。")

heading(doc, "5.4 Case D：不支持或服务失败", 2)
add_table(doc, ["场景", "系统行为", "用户文案原则"], [
    ["格式不是 Image / Poster", "返回 Outside scope", "说明当前组合没有对应知识包。"],
    ["OCR 无法形成可靠文本", "返回 Need input 或 Error", "要求确认文字或转人工，不假装已检查。"],
    ["知识包缺失/过期", "禁止调用通用模型生成规则", "明确覆盖边界和替代路径。"],
    ["模型或规则服务超时", "旧结果失效，阻止关键推进，允许重试", "透明说明暂时不可用。"],
], [2200, 3500, 3660], font_size=9.1)

add_kicker(doc, "06 · KNOWLEDGE OPERATIONS")
heading(doc, "6. 合规知识治理", 1)
add_table(doc, ["模块", "核心能力", "关键权限"], [
    ["知识包", "Country × Industry × Product × Regulation × Format；版本、有效期与回滚", "知识管理员"],
    ["来源与规则", "官方来源快照、规则映射、结构化条件与结论", "法规专家"],
    ["内部质量复核", "模型错误样本、OCR 失败、抽样复核和错误分类", "审核员/法务"],
    ["评测与发布", "回归集、阈值、严重漏放门禁、版本发布与回滚", "发布管理员"],
], [1800, 5600, 1960], font_size=9.1)

add_kicker(doc, "07 · FUNCTIONAL REQUIREMENTS")
heading(doc, "7. P0 功能清单", 1)
add_table(doc, ["ID", "功能", "P0 验收标准"], [
    ["F01", "海报生成", "输入 Prompt 后生成新的背景图和可编辑文字层；不得继续展示原示例内容。"],
    ["F02", "触发事件", "仅在海报渲染完成、文字/Context 变化、应用改写或继续使用前触发。"],
    ["F03", "文字提取", "原生文字图层优先，OCR 兜底；输出区域、来源和置信度。"],
    ["F04", "Context 采集", "Country、Industry、Product、Product Regulation、Channel、Format 可追溯；缺失时追问。"],
    ["F05", "知识包路由", "仅命中已发布且有效的 AU-BEA-POSTER-TEXT 版本。"],
    ["F06", "四态决策", "输出 Ready、Fix、Need input、Outside scope；系统失败单独返回 Error。"],
    ["F07", "问题定位", "返回海报文字区域、问题类型、原因、可能规则与下一步。"],
    ["F08", "最小改写", "只修改必要文字层；写回后自动重检，修复本身不等于通过。"],
    ["F09", "覆盖披露", "所有结果展示 checked/not checked、知识版本和 trace_id。"],
    ["F10", "发布 Gate", "非有效 Ready 不得下载、发布或继续；不阻断对话和编辑。"],
    ["F11", "缓存与失效", "内容哈希 + Context + Rule Version 一致时复用；任一变化立即失效。"],
    ["F12", "安全与密钥", "图片与语言模型 Key 仅存服务端 Secret；前端和日志不得泄露。"],
], [780, 2100, 6480], font_size=8.7)

add_kicker(doc, "08 · SYSTEM & DATA")
heading(doc, "8. 系统架构与决策契约", 1)
heading(doc, "8.1 模型与规则分工", 2)
add_table(doc, ["组件", "职责", "禁止事项"], [
    ["CogView / 图像模型", "根据 Prompt 生成广告背景图，尽量不烘焙文字", "不输出合规结论"],
    ["GLM 文案模型", "生成品牌名、短标题和可编辑海报文案", "不决定法规是否适用"],
    ["文字提取层", "读取 Canva 文字图层；对扁平化图片执行 OCR", "不把低置信文本当确定输入"],
    ["知识路由与规则层", "决定覆盖范围、强制声明、规则版本和发布 Gate", "不存在知识包时不得让模型补齐"],
    ["GLM 合规推理", "语义理解、片段定位、受控解释与最小改写", "不得发明法律、规则 ID 或来源"],
], [1900, 4200, 3260], font_size=8.9)

heading(doc, "8.2 决策输出结构", 2)
add_table(doc, ["字段", "示例/规则", "用途"], [
    ["status", "READY / FIX / NEED_INPUT / OUTSIDE_SCOPE / ERROR", "统一前端状态机"],
    ["extraction", "source、regions[]、confidence、text_hash", "证明审核对象来自最终海报"],
    ["issues[]", "region_id、span、issue_type、message、suggestion", "区域定位与修复"],
    ["evidence[]", "source_id、rule_id、excerpt_hash、effective_at", "解释与审计"],
    ["coverage", "checked、not_checked、jurisdiction、format", "边界披露"],
    ["knowledge_version", "AU-BEA-POSTER-TEXT v1.4", "重现与回滚"],
    ["models", "copy_model、image_model、decision_model", "版本追踪"],
    ["trace_id / latency", "唯一链路 ID / 分阶段耗时", "观测与故障定位"],
], [1800, 4300, 3260], font_size=8.9)

heading(doc, "8.3 隐私与日志", 2)
add_bullet(doc, "默认不使用客户 Prompt、图片或原文训练模型；需明确授权并去标识后才能进入评测集。")
add_bullet(doc, "生产日志优先保存内容哈希、区域坐标、规则 ID、版本和状态；原文按最短必要保留期处理。")
add_bullet(doc, "模型 Key、系统 Prompt 和内部来源快照不得发送到浏览器端。")
add_bullet(doc, "临时图片 URL 有有效期；生产方案需转存到受控存储并执行访问控制。")

add_kicker(doc, "09 · EVALUATION")
heading(doc, "9. 评测体系与发布门禁", 1)
add_table(doc, ["维度", "核心指标", "发布门禁"], [
    ["生成有效性", "输入生效率、背景图替换率、文字层非空率", "不得继续显示示例内容；失败需明确提示"],
    ["文字提取", "字符准确率、区域召回、低置信识别率", "强制声明区域不得因 OCR 漏掉"],
    ["安全性", "关键遗漏召回、严重违规漏放率", "严重漏放为 0；召回达到知识包阈值"],
    ["可用性", "误拦截率、修复后完成率", "不劣于当前生产版本"],
    ["边界", "缺 Context / 不支持组合识别率", "不得输出 Ready"],
    ["可解释性", "证据可追溯率、错误 rule_id 率", "证据覆盖 100%；无错误映射"],
    ["稳定性", "P95 延迟、超时率、版本一致性", "满足 SLO；支持回滚"],
], [1700, 4200, 3460], font_size=8.9)
heading(doc, "9.1 必测样本类型", 2)
add_bullet(doc, "Mandatory disclosure：AUST L 海报缺少必要声明，但没有任何违禁词。")
add_bullet(doc, "Regulatory boundary：Cosmetic 海报出现 repair damaged skin 等治疗边界声称。")
add_bullet(doc, "Evidence dependency：量化/临床声称但证据与产品、时间范围不匹配。")
add_bullet(doc, "Hard negatives：主观偏好、成分事实、非量化肤感等不应被过度拦截。")
add_bullet(doc, "OCR stress：小字号、低对比、弯曲文字、复杂背景和多语言文字。")

add_kicker(doc, "10 · ROLLOUT & RISK")
heading(doc, "10. 上线计划、风险与依赖", 1)
heading(doc, "10.1 上线阶段", 2)
add_table(doc, ["阶段", "范围", "退出条件"], [
    ["Shadow", "生成后运行检查但不阻止用户", "完成基线、关键遗漏与误拦截校准"],
    ["Assisted", "展示四态与建议；仅对高置信强制义务硬阻断", "专家复核稳定，用户理解边界"],
    ["Controlled GA", "白名单 AU Beauty 用户与知识包版本", "指标达标、回滚演练完成"],
    ["Expansion", "更多 AU 品类、OCR 上传图、平台政策包", "每个包独立通过评测门禁"],
], [1700, 4300, 3360], font_size=9.0)

heading(doc, "10.2 风险清单", 2)
add_table(doc, ["风险", "影响", "控制措施", "Owner"], [
    ["图片生成失败或过慢", "用户误以为输入无效", "进度态、超时提示、文字层兜底；不得静默保留示例", "Product + ML"],
    ["图片内烘焙文字未被读取", "漏检", "生成时要求无文字；上传/扁平图使用 OCR；展示提取来源", "ML"],
    ["OCR 漏掉小字声明", "严重漏放", "区域检测、低置信门禁、强制区域校验、人工抽样", "Risk + ML"],
    ["误拦截过多", "发布转化下降", "Hard negative 集、可解释结果、灰度策略", "Product"],
    ["规则过期", "错误结论", "来源 last-checked、到期告警、紧急下线/回滚", "Knowledge Ops"],
    ["覆盖误解", "局部结论被当作全局", "固定展示 checked/not checked；Ready 限定措辞", "Design"],
    ["模型 Key 泄露", "费用与安全风险", "服务端 Secret、速率限制、日志脱敏与轮换", "Security"],
], [1800, 2100, 4060, 1400], font_size=8.6)

add_kicker(doc, "11 · ACCEPTANCE")
heading(doc, "11. P0 端到端验收", 1)
add_table(doc, ["场景", "Given / When", "Then"], [
    ["输入生成", "用户输入新的广告海报需求并发送", "生成新的背景图和文字层；示例内容被替换；完成后自动检查。"],
    ["缺少必填声明", "AUST L 海报缺 mandatory statement", "返回 Fix；定位缺失区域；补充后重检才允许继续。"],
    ["分类未知", "Product Regulation = Not sure，海报含功效声称", "返回 Need input；不得从 Prompt 推断分类。"],
    ["检查通过", "AU Cosmetic 海报文字无明确问题", "显示 Ready、Checked and fine、覆盖范围、版本和 trace_id。"],
    ["文字编辑", "用户修改海报文字层并停止编辑", "旧结果立即失效；debounce 后自动重检。"],
    ["应用改写", "用户点击补充声明/建议改写", "写回文字层并重检；修复动作本身不等于通过。"],
    ["范围不支持", "非 AU Beauty Image / Poster 组合", "返回 Outside scope；不显示 Ready。"],
    ["系统失败", "生成、OCR、模型或知识包不可用", "显示 Error/不可用；阻止关键推进；提供重试。"],
], [1900, 3550, 3910], font_size=8.8)

heading(doc, "11.1 立项前待决问题", 2)
add_table(doc, ["问题", "建议决策"], [
    ["P0 是否硬阻断发布？", "仅对高置信强制义务硬阻断；其余根据知识包策略进入 Need input 或警示。"],
    ["OCR 置信度阈值如何设置？", "按字号、区域与义务严重度分层；关键声明区域采用更严格阈值。"],
    ["图片生成模型如何选择？", "Demo 使用 CogView 快速生成；生产按质量、延迟、成本与数据驻留评估。"],
    ["谁拥有规则最终解释权？", "法规专家负责规则，产品负责体验，发布管理员负责上线门禁。"],
    ["哪些内容可用于模型改进？", "默认不使用客户原文训练；仅在授权、去标识和保留期控制后进入数据集。"],
    ["Ready 如何表述？", "固定使用“在已识别并检查的海报文字范围内”，并展示 not checked 列表。"],
], [3000, 6360], font_size=9.0)

add_kicker(doc, "APPENDIX")
heading(doc, "附录 A：Demo 展示管理员模式", 1)
add_body(doc, "展示管理员模式仅用于评审时比较“页面内”和“侧边抽屉”两种呈现方式，不属于用户权限系统，也不进入正式产品流程。300ms 内双击 Ctrl（Mac 为 Cmd）可调出设置；再次双击或按 Esc 关闭，选择保存在 localStorage。")
add_image(doc, "04-admin-display-mode.png", "图 6  展示管理员模式：同一份状态与数据，只切换布局容器。")
add_bullet(doc, "页面内：信息完整，适合详细审阅；结果紧贴生成海报下沿。")
add_bullet(doc, "侧边：400px 右侧抽屉，主内容同步压窄；用户始终看得到海报。")
add_bullet(doc, "切换展示形式不重新运行检查，复用当前结果并保持定位。")

heading(doc, "附录 B：术语", 1)
add_table(doc, ["术语", "定义"], [
    ["Usable marketing content", "已经形成、可被选择、编辑、下载、发布或继续使用的营销图片/海报。"],
    ["Visible poster text", "最终画面中用户可见的品牌名、标题、正文、价格、声明、脚注等文字。"],
    ["Text layer", "Canva 内可编辑的结构化文字对象，优先于 OCR。"],
    ["OCR fallback", "对上传或扁平化图片中的文字进行识别，并返回区域与置信度。"],
    ["Knowledge package", "按国家、行业、产品、监管身份和内容格式发布的规则与来源集合。"],
    ["Pre-use Compliance Gate", "在下载、发布或继续之前确认当前结果存在、有效且允许推进。"],
], [2200, 7160], font_size=9.1)

doc.core_properties.title = "Canva Grow 营销内容合规层 PRD（更新版）"
doc.core_properties.subject = "Australia Beauty Image/Poster Visible Text Compliance"
doc.core_properties.author = "AI Product · Risk & Compliance"
doc.core_properties.keywords = "Canva, Compliance Layer, Poster, OCR, Australia, Beauty"
doc.save(OUT)
print(OUT)
